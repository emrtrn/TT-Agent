"""TT-Agent Markdown kaynakları için DOCX dışa aktarıcı."""
from __future__ import annotations

import argparse
import hashlib
import html
import json
import re
import sys
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
UNITS = ROOT / "units"
DEFAULT_OUTPUT = ROOT / "docs" / "docx"
THEME = Path(__file__).with_name("docx-theme.json")
BREAK = re.compile(r"^\s*<!--\s*(?:PAGEBREAK|PDF_PAGE_BREAK)\s*-->\s*$", re.I)
DIVIDER = re.compile(r"^\|?\s*:?-{3,}:?\s*(?:\|\s*:?-{3,}:?\s*)+\|?$")
WRITE_LINE = re.compile(r"^(?:\\_){3,}\s*$")
WRITE_FIELD = re.compile(r"(?:\\_){3,}")
INLINE = re.compile(
    r"(?P<break><br\s*/?>)|"
    r"(?P<field>(?:\\_){3,})|"
    r"(?P<bold>\*\*(?P<bold_text>.+?)\*\*)|"
    r"(?P<link>\[(?P<link_text>[^\]]+)\]\((?P<link_url>[^)]+)\))",
    re.I,
)


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def heading(markdown: str, fallback: str) -> str:
    found = re.search(r"^#{1,6}\s+(.+?)\s*$", markdown, re.M)
    return found.group(1).strip() if found else fallback


def unit_name(path: Path) -> str:
    found = re.search(r"unit(\d+)", path.as_posix(), re.I)
    return f"Ünite {found.group(1)}" if found else "Ünite"


def markdown_tables(markdown: str) -> int:
    lines = markdown.splitlines()
    return sum("|" in line and bool(DIVIDER.match(lines[index + 1].strip())) for index, line in enumerate(lines[:-1]))


def markdown_breaks(markdown: str) -> int:
    return sum(bool(BREAK.match(line)) for line in markdown.splitlines())


def page_number(paragraph) -> None:
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run, text = OxmlElement("w:r"), OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    field.append(run)
    paragraph._p.append(field)


def cell_style(cell, fill: str, border: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    props.append(shade)
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "4")
        item.set(qn("w:color"), border)
        borders.append(item)
    props.append(borders)


class DocxRenderer:
    def __init__(self, source: Path, markdown: str, theme: dict) -> None:
        self.source, self.markdown, self.theme, self.colors = source, markdown, theme, theme["colors"]
        self.doc = Document()
        self.configure()

    def add_run(self, paragraph, text: str, size: float | None = None, bold: bool = False, color: str = "text") -> None:
        run = paragraph.add_run(text)
        run.font.name = self.theme["document"]["font"]
        run._element.rPr.rFonts.set(qn("w:eastAsia"), self.theme["document"]["font"])
        run.font.size = Pt(size or self.theme["document"]["bodyFontSizePt"])
        run.font.color.rgb = RGBColor.from_string(self.colors[color])
        run.bold = bold

    def add_hyperlink(self, paragraph, text: str, url: str, size: float | None = None, bold: bool = False) -> None:
        relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), relationship_id)
        run, properties, font = OxmlElement("w:r"), OxmlElement("w:rPr"), OxmlElement("w:rFonts")
        font.set(qn("w:ascii"), self.theme["document"]["font"])
        font.set(qn("w:hAnsi"), self.theme["document"]["font"])
        font.set(qn("w:eastAsia"), self.theme["document"]["font"])
        properties.append(font)
        font_size = OxmlElement("w:sz")
        font_size.set(qn("w:val"), str(round((size or self.theme["document"]["bodyFontSizePt"]) * 2)))
        properties.append(font_size)
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "0563C1")
        properties.append(color)
        underline = OxmlElement("w:u")
        underline.set(qn("w:val"), "single")
        properties.append(underline)
        if bold:
            properties.append(OxmlElement("w:b"))
        run.append(properties)
        value = OxmlElement("w:t")
        value.text = text
        run.append(value)
        hyperlink.append(run)
        paragraph._p.append(hyperlink)

    def add_text(self, paragraph, text: str, size: float | None = None, bold: bool = False, color: str = "text") -> None:
        text = html.unescape(text)
        cursor = 0
        for match in INLINE.finditer(text):
            if match.start() > cursor:
                self.add_run(paragraph, text[cursor:match.start()], size, bold, color)
            if match.group("break"):
                self.add_run(paragraph, "", size, bold, color)
                paragraph.runs[-1].add_break()
            elif match.group("field"):
                self.add_run(paragraph, "\u00A0" * match.group("field").count("\\_"), size, bold, color)
                paragraph.runs[-1].font.underline = True
            elif match.group("bold"):
                self.add_text(paragraph, match.group("bold_text"), size, True, color)
            elif match.group("link"):
                self.add_hyperlink(paragraph, match.group("link_text"), match.group("link_url"), size, bold)
            cursor = match.end()
        if cursor < len(text):
            self.add_run(paragraph, text[cursor:], size, bold, color)

    def configure(self) -> None:
        page, doc = self.theme["page"], self.theme["document"]
        section = self.doc.sections[0]
        section.page_width, section.page_height = Cm(page["widthCm"]), Cm(page["heightCm"])
        section.top_margin, section.bottom_margin = Cm(page["topMarginCm"]), Cm(page["bottomMarginCm"])
        section.left_margin, section.right_margin = Cm(page["leftMarginCm"]), Cm(page["rightMarginCm"])
        section.header_distance, section.footer_distance = Cm(page["headerDistanceCm"]), Cm(page["footerDistanceCm"])
        normal = self.doc.styles["Normal"]
        normal.font.name, normal.font.size = doc["font"], Pt(doc["bodyFontSizePt"])
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), doc["font"])
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        self.add_run(header, f'{self.theme["identity"]["header"]}  |  {unit_name(self.source)}', 8.5, True, "muted")
        header._p.get_or_add_pPr().append(parse_xml(r'<w:pBdr {}><w:bottom w:val="single" w:sz="6" w:space="3" w:color="{}"/></w:pBdr>'.format(nsdecls("w"), self.colors["headerRule"])))
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.add_run(footer, f'{self.theme["identity"]["footer"]}  |  Sayfa ', 8.5, color="muted")
        page_number(footer)

    def add_table(self, rows: list[list[str]]) -> None:
        columns = max(len(row) for row in rows)
        table = self.doc.add_table(rows=0, cols=columns)
        for row_index, values in enumerate(rows):
            row, is_header = table.add_row(), row_index == 0
            props = row._tr.get_or_add_trPr()
            props.append(OxmlElement("w:cantSplit"))
            if is_header:
                marker = OxmlElement("w:tblHeader")
                marker.set(qn("w:val"), "true")
                props.append(marker)
            for index, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                fill = self.colors["primary"] if is_header else self.colors["tableStripe"] if row_index % 2 == 0 else "FFFFFF"
                cell_style(cell, fill, self.colors["grid"])
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                self.add_text(p, values[index].strip() if index < len(values) else "", self.theme["document"]["tableFontSizePt"], is_header, "white" if is_header else "text")
        self.doc.add_paragraph()

    def add_writing_line(self) -> None:
        p = self.doc.add_paragraph()
        self.add_run(p, " ")
        p._p.get_or_add_pPr().append(parse_xml(r'<w:pBdr {}><w:bottom w:val="single" w:sz="8" w:space="2" w:color="{}"/></w:pBdr>'.format(nsdecls("w"), self.colors["grid"])))

    def render(self) -> Document:
        title = heading(self.markdown, self.source.stem)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.add_run(p, title, self.theme["document"]["titleFontSizePt"], True, "primary")
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.add_run(p, f"8. Sınıf | Teknoloji ve Tasarım | {unit_name(self.source)}", 10.5, True, "secondary")
        lines, index, first_heading = self.markdown.splitlines(), 0, True
        while index < len(lines):
            line, stripped = lines[index], lines[index].strip()
            if not stripped:
                index += 1
            elif BREAK.match(line):
                self.doc.add_page_break()
                index += 1
            elif WRITE_LINE.fullmatch(stripped):
                self.add_writing_line()
                index += 1
            elif index + 1 < len(lines) and "|" in line and DIVIDER.match(lines[index + 1].strip()):
                rows, index = [line.strip().strip("|").split("|")], index + 2
                while index < len(lines) and lines[index].strip().startswith("|"):
                    rows.append(lines[index].strip().strip("|").split("|"))
                    index += 1
                self.add_table(rows)
            elif match := re.match(r"^(#{1,6})\s+(.+?)\s*$", line):
                if not first_heading:
                    p = self.doc.add_paragraph()
                    p.paragraph_format.keep_with_next = True
                    level = len(match.group(1))
                    size, color = (self.theme["document"]["majorHeadingFontSizePt"], "primary") if level == 1 else (self.theme["document"]["subHeadingFontSizePt"], "accent")
                    self.add_text(p, match.group(2).strip(), size, True, color)
                first_heading = False
                index += 1
            elif item := re.match(r"^\s*([-+*]|\d+[.)])\s+(.+)$", line):
                checkbox = re.match(r"^\[([xX ])\]\s*(.*)$", item.group(2))
                if checkbox:
                    p = self.doc.add_paragraph()
                    self.add_text(p, ("■ " if checkbox.group(1).lower() == "x" else "□ ") + checkbox.group(2))
                else:
                    p = self.doc.add_paragraph(style="List Number" if item.group(1)[0].isdigit() else "List Bullet")
                    self.add_text(p, item.group(2))
                index += 1
            elif stripped.startswith(">"):
                p = self.doc.add_paragraph()
                p._p.get_or_add_pPr().append(parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls("w"), self.colors["veryPale"])))
                self.add_text(p, re.sub(r"^\s*>\s?", "", line))
                index += 1
            else:
                p = self.doc.add_paragraph()
                self.add_text(p, line)
                index += 1
        self.doc.core_properties.title = title
        return self.doc


def exported_text(document: Document) -> str:
    pieces = [p.text for p in document.paragraphs]
    pieces.extend(p.text for table in document.tables for row in table.rows for cell in row.cells for p in cell.paragraphs)
    for section in document.sections:
        pieces.extend(p.text for p in section.header.paragraphs + section.footer.paragraphs)
    return "\n".join(pieces)


def validate(source: Path, output: Path, markdown: str, before: str) -> dict:
    if sha(source) != before:
        raise RuntimeError(f"Kaynak Markdown üretim sırasında değişti: {source}")
    with ZipFile(output) as package:
        required = {"[Content_Types].xml", "word/document.xml", "word/header1.xml", "word/footer1.xml"}
        if required - set(package.namelist()):
            raise RuntimeError(f"DOCX paketi eksik: {output}")
        if package.read("word/document.xml").count(b'<w:br w:type="page"/>') != markdown_breaks(markdown):
            raise RuntimeError(f"Sayfa sonu doğrulaması başarısız: {output}")
        if b"PAGE" not in package.read("word/footer1.xml"):
            raise RuntimeError(f"Sayfa numarası doğrulaması başarısız: {output}")
    doc = Document(output)
    if len(doc.tables) != markdown_tables(markdown):
        raise RuntimeError(f"Tablo doğrulaması başarısız: {output}")
    if not doc.paragraphs or doc.paragraphs[0].text != heading(markdown, source.stem):
        raise RuntimeError(f"Başlık doğrulaması başarısız: {output}")
    expected = {char for char in markdown if char in "ÇçĞğİıÖöŞşÜü"}
    if not expected.issubset(set(exported_text(doc))):
        raise RuntimeError(f"Türkçe karakter doğrulaması başarısız: {output}")
    rendered = exported_text(doc)
    for raw in ("**", "<br", "&nbsp;"):
        if raw in markdown and raw in rendered:
            raise RuntimeError(f"Satır içi Markdown/HTML işleme doğrulaması başarısız: {output}")
    return {"source": source, "output": output, "tables": markdown_tables(markdown), "breaks": markdown_breaks(markdown)}


def selected_sources(grade: str, selected_unit: str | None, all_units: bool, input_path: Path | None) -> list[Path]:
    if input_path:
        return [input_path.resolve()]
    root = UNITS / f"{grade}_sinif"
    names = sorted(path.name for path in root.glob("unit*") if path.is_dir()) if all_units else [selected_unit]
    files = [file for name in names for file in sorted((root / str(name)).glob("*.md"))]
    if not files:
        raise RuntimeError("Seçimde dışa aktarılacak Markdown bulunamadı.")
    return files


def main() -> None:
    parser = argparse.ArgumentParser()
    choice = parser.add_mutually_exclusive_group(required=True)
    choice.add_argument("source", nargs="?", type=Path)
    choice.add_argument("--unit")
    choice.add_argument("--all", action="store_true")
    parser.add_argument("--grade", default="8")
    parser.add_argument("--output-root", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    theme, output_root = json.loads(THEME.read_text(encoding="utf-8")), args.output_root.resolve()
    results = []
    for source in selected_sources(args.grade, args.unit, args.all, args.source):
        before, markdown = sha(source), source.read_text(encoding="utf-8")
        output = output_root / source.parent.relative_to(UNITS) / f"{source.stem}.docx"
        output.parent.mkdir(parents=True, exist_ok=True)
        DocxRenderer(source, markdown, theme).render().save(output)
        results.append(validate(source, output, markdown, before))
    for result in results:
        print(f'{result["source"].relative_to(ROOT)} -> {result["output"].relative_to(ROOT)} | {result["tables"]} tablo | {result["breaks"]} sayfa sonu')
    print(f"Tamamlandı: {len(results)} DOCX")


if __name__ == "__main__":
    try:
        main()
    except (OSError, RuntimeError, ValueError) as error:
        print(f"HATA: {error}", file=sys.stderr)
        raise SystemExit(1) from error
